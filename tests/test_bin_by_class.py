# 評價清單 #2：判到的 bin × 真正的類別（ground truth 標了 type 才有）。
from __future__ import annotations

from d4t.core.export.report import UNBINNED_KEY, summarize, write_excel

ROWS = [
    {"defect_id": "1", "ok": True, "bin": 1, "score": 9.0, "features": {}},
    {"defect_id": "2", "ok": True, "bin": 1, "score": 8.0, "features": {}},
    {"defect_id": "3", "ok": True, "bin": 0, "score": 1.0, "features": {}},
    {"defect_id": "4", "ok": True, "bin": 0, "score": 1.0, "features": {}},
    {"defect_id": "5", "ok": False, "bin": None, "score": None, "features": {}},
]
GT = {"1": {"is_real": True, "type": "bridge"},
      "2": {"is_real": True, "type": "dark_blob"},
      "3": {"is_real": True, "type": "dark_blob"},   # 漏抓的那一顆是哪一類
      "4": {"is_real": False, "type": "none"}}


def test_the_table_says_which_class_landed_where():
    t = summarize(ROWS, ground_truth=GT)["bin_by_class"]
    assert t["classes"] == ["bridge", "dark_blob", "none"]
    rows = {r["bin"]: r for r in t["rows"]}
    assert rows["1"]["by_class"] == {"bridge": 1, "dark_blob": 1}
    assert rows["0"]["by_class"] == {"dark_blob": 1, "none": 1}
    assert rows[UNBINNED_KEY]["unlabelled"] == 1
    assert t["rows"][-1]["bin"] == UNBINNED_KEY, "沒判定的那一格排最後"
    assert [r["bin"] for r in t["rows"]][:2] == ["1", "0"], "bin 大到小"


def test_no_classes_no_table():
    """只標了真／假 → 不印這張表（全空的表比沒有更像壞了）。"""
    s = summarize(ROWS, ground_truth={"1": True, "3": False})
    assert "bin_by_class" not in s
    assert "bin_by_class" not in summarize(ROWS)


def test_the_workbook_has_the_table(tmp_path):
    from openpyxl import load_workbook
    path = write_excel(ROWS, str(tmp_path / "r.xlsx"), ground_truth=GT)
    ws = load_workbook(path).worksheets[0]
    cells = [[c.value for c in row] for row in ws.iter_rows()]
    flat = [v for row in cells for v in row if v is not None]
    assert "Bin × actual class" in flat
    hdr = next(r for r in cells if r[0] == "Called")
    assert "Actual: dark_blob" in hdr
